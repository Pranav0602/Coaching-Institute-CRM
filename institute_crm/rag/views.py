"""
HTTP controller layer for the RAG app.
"""
from __future__ import annotations

from rest_framework import viewsets, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import AllowAny, IsAuthenticated

from rag.models import KnowledgeDocument, IngestionJob, RagQueryAudit
from rag.serializers import (
    KnowledgeDocumentSerializer, RagQueryInputSerializer, 
    RagQueryAuditSerializer, IngestionJobSerializer
)
from rag.services import IngestionService, GenerationService
from accounts.permissions import IsBranchAdmin, IsAdminOrCounselorReadOnly
from accounts.models import Role
from django.db.models import Q


class RagQueryView(APIView):
    """
    ``POST /api/v1/rag/query/``
    Public & Authenticated knowledge assistant endpoint.
    """
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = RagQueryInputSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        user = request.user if request.user.is_authenticated else None
        result = GenerationService.answer_query(
            prompt=serializer.validated_data['query'],
            actor=user,
            session_id=serializer.validated_data.get('session_id')
        )
        return Response(result, status=status.HTTP_200_OK)


class IngestionTriggerView(APIView):
    """
    ``POST /api/v1/rag/ingest/``
    Triggers catalog and FAQ knowledge base indexing.
    """
    permission_classes = [IsAuthenticated, IsBranchAdmin]

    def post(self, request):
        job = IngestionService.sync_academic_catalogue()
        return Response(IngestionJobSerializer(job).data, status=status.HTTP_202_ACCEPTED)


class KnowledgeDocumentViewSet(viewsets.ModelViewSet):
    """
    ``/api/v1/rag/documents/``
    CRUD for SUPER_ADMIN / BRANCH_ADMIN; read-only published docs for
    ADMISSION_COUNSELOR; retrieval itself is via the public /rag/query endpoint.
    """
    queryset = KnowledgeDocument.objects.filter(is_deleted=False).prefetch_related('chunks')
    serializer_class = KnowledgeDocumentSerializer
    permission_classes = [IsAuthenticated, IsAdminOrCounselorReadOnly]

    def get_queryset(self):
        user = self.request.user
        qs = KnowledgeDocument.objects.filter(is_deleted=False).prefetch_related('chunks')
        role_code = getattr(getattr(user, "role", None), "code", None)
        is_admin = bool(
            getattr(user, "is_superuser", False)
            or role_code in [Role.SUPER_ADMIN, Role.BRANCH_ADMIN]
        )
        # Counselors (and any other read-only caller) only see published docs.
        if not is_admin:
            qs = qs.filter(is_published=True)

        params = self.request.query_params
        category = params.get("category")
        if category:
            qs = qs.filter(category=category)
        course_id = params.get("course_id")
        if course_id:
            qs = qs.filter(
                Q(metadata_json__course_id=str(course_id))
                | Q(metadata_json__courseId=str(course_id))
            )
        search = params.get("search")
        if search:
            qs = qs.filter(Q(title__icontains=search) | Q(content__icontains=search))
        return qs.order_by("-updated_at")

    def perform_create(self, serializer):
        doc = serializer.save()
        IngestionService.index_document(doc)

    def perform_update(self, serializer):
        doc = serializer.save()
        IngestionService.index_document(doc)


class DocumentTextExtractionView(APIView):
    """
    ``POST /api/v1/rag/extract-text/``
    Accepts a document file (PDF, DOCX, TXT, MD, CSV) and returns its extracted text.
    """
    permission_classes = [IsAuthenticated, IsAdminOrCounselorReadOnly]
    MAX_FILE_BYTES = 10 * 1024 * 1024

    def post(self, request):
        upload = request.FILES.get('file')
        if not upload:
            return Response({'detail': 'No file provided (multipart field "file").'}, status=status.HTTP_400_BAD_REQUEST)
        if upload.size > self.MAX_FILE_BYTES:
            return Response({'detail': 'File too large (max 10 MB).'}, status=status.HTTP_400_BAD_REQUEST)

        name = (upload.name or '').lower()
        try:
            if name.endswith('.pdf'):
                text = self._extract_pdf(upload)
            elif name.endswith('.docx'):
                text = self._extract_docx(upload)
            elif name.endswith(('.txt', '.md', '.csv')):
                text = upload.read().decode('utf-8', errors='replace')
            else:
                return Response({'detail': 'Unsupported file type. Use .pdf, .docx, .txt, .md or .csv.'}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:
            return Response({'detail': f'Could not extract text: {exc}'}, status=status.HTTP_422_UNPROCESSABLE_ENTITY)

        text = (text or '').strip()
        if not text:
            return Response({'detail': 'No extractable text found (scanned/image PDFs are not supported).'}, status=status.HTTP_422_UNPROCESSABLE_ENTITY)

        file_url = self._archive_to_s3(upload)
        return Response({'file_name': upload.name, 'text': text, 'file_url': file_url}, status=status.HTTP_200_OK)

    @staticmethod
    def _archive_to_s3(upload) -> str | None:
        """Best-effort archive of the original file to Supabase S3; text extraction must not fail if storage is unavailable."""
        try:
            import uuid
            from aws_services.s3_service import s3_service
            key = f"knowledge_base/{uuid.uuid4().hex}_{upload.name}"
            return s3_service.upload_fileobj(upload, key, content_type=upload.content_type)
        except Exception:
            return None

    @staticmethod
    def _extract_pdf(upload) -> str:
        from pypdf import PdfReader
        reader = PdfReader(upload)
        return "\n".join((page.extract_text() or '') for page in reader.pages)

    @staticmethod
    def _extract_docx(upload) -> str:
        import docx
        document = docx.Document(upload)
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))
        return "\n".join(parts)


class RagStatsView(APIView):
    """
    ``GET /api/v1/rag/stats/``
    Provides summary statistics of the RAG vector corpus.
    """
    permission_classes = [IsAuthenticated, IsBranchAdmin]

    def get(self, request):
        from rag.models import DocumentChunk
        doc_count = KnowledgeDocument.objects.filter(is_deleted=False).count()
        chunk_count = DocumentChunk.objects.filter(is_deleted=False).count()
        queries_count = RagQueryAudit.objects.count()

        return Response({
            "documents_count": doc_count,
            "chunks_count": chunk_count,
            "queries_served": queries_count,
            "status": "ready"
        })
